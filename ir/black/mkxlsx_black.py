import os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
HERE=os.path.dirname(os.path.abspath(__file__)); F="Arial"
wb=Workbook(); ws=wb.active; ws.title="Состав"
thin=Side(style="thin",color="A6A6A6"); B=Border(top=thin,bottom=thin,left=thin,right=thin)
hdr=PatternFill("solid",fgColor="1F4E79"); sfill=PatternFill("solid",fgColor="DDEBF7")
ws["A1"]="Состав органической части образца buryak_shmurdyak_black (в пересчёте на 100 % органики)"; ws["A1"].font=Font(name=F,bold=True,size=13)
ws["A2"]="Оценка по ИК-спектроскопии (НПВО). ЯМР будет добавлен после получения спектров. Неорганический компонент в 100 % не входит."; ws["A2"].font=Font(name=F,italic=True,size=9,color="404040")
H=["Образец","Компонент","ИК, %","ЯМР, %","Среднее по ИК и ЯМР, %","Итог, % (сумма = 100)","Диапазон, %","Идентификация и комментарий","Неорганическая часть (не входит в 100 %)"]
r0=4
for j,h in enumerate(H,1):
    c=ws.cell(r0,j,h); c.font=Font(name=F,bold=True,color="FFFFFF",size=10); c.fill=hdr; c.border=B
    c.alignment=Alignment(wrap_text=True,vertical="center",horizontal="center")
rows=[("Дибутилфталат (ДБФ)",100,"ожидается","95–100","Единственный обнаруженный органический компонент. ИК: C=O 1723 см⁻¹; дублет ароматики 1600/1580 см⁻¹; 1277, 1119, 1072 см⁻¹ (сложный эфир фталевой кислоты); 742 см⁻¹ (орто-замещённое кольцо); н-бутил — 2960, 1384, 962/941 см⁻¹. Длинноцепочечные фталаты (ДОФ) не подходят: нет сильной полосы 2858 см⁻¹. Отличить от диизобутилфталата окончательно — по ЯМР ¹H (4,30 т против 4,08 д м.д.)."),
 ("Прочие органические добавки (стеариновая кислота и др.)","не обнаружено (< 2–3 %)","ожидается","< 3","В типовых вулканизующих пастах ≈1 % стеариновой кислоты — ниже предела обнаружения ИК."),
 ("Тиокол","не обнаружен (< 5 %)","ожидается","< 5","Нет полос формали 1024/949 см⁻¹, S–CH₂ 1410 см⁻¹ и SH 2560 см⁻¹. Образец — не основа герметика."),
 ("Эпоксидная смола","не обнаружена","ожидается","—","Нет полос бисфенола А 1510, 1244, 831 см⁻¹."),
 ("Хлорпарафин","не обнаружен","ожидается","—","Нет отдельной полосы CHCl 1260–1270 см⁻¹ и прироста CH₂. Окончательно — по ЯМР ¹H (CHCl 3,8–4,5 м.д.)."),
 ("Полиуретан / лапрол","не обнаружены","ожидается","—","Нет полос N–H и широкой полосы простого эфира.")]
r=r0+1; start=r
for comp,ir,nmr,rng,cmt in rows:
    ws.cell(r,2,comp); ws.cell(r,3,ir); ws.cell(r,4,nmr)
    ws.cell(r,5,f'=IF(COUNT(C{r}:D{r})=0,"—",AVERAGE(C{r}:D{r}))')
    ws.cell(r,7,rng); ws.cell(r,8,cmt); r+=1
end=r-1
for rr in range(start,end+1):
    ws.cell(rr,6,f'=IF(ISNUMBER(E{rr}),E{rr}/SUM($E${start}:$E${end}),"—")')
ws.cell(start,1,"buryak_shmurdyak_black")
ws.cell(start,9,"Чёрный неорганический компонент: рост поглощения M–O ниже 650 см⁻¹; предположительно диоксид марганца MnO₂. Количество — по зольности, металл — РФА.")
ws.merge_cells(start_row=start,end_row=end,start_column=1,end_column=1); ws.merge_cells(start_row=start,end_row=end,start_column=9,end_column=9)
ws.cell(r,2,"Сумма органической части"); ws.cell(r,6,f"=SUM(F{start}:F{end})")
for j in range(1,10):
    c=ws.cell(r,j); c.fill=sfill; c.font=Font(name=F,bold=True,size=9); c.border=B
ws.cell(r,6).number_format="0.0%"
for row in ws.iter_rows(min_row=start,max_row=end):
    for c in row:
        c.font=Font(name=F,size=9,bold=(c.column==1)); c.border=B
        c.alignment=Alignment(wrap_text=True,vertical="top",horizontal="center" if c.column in (3,4,5,6,7) else "left")
        if c.column in (3,5): c.number_format="0.0"
        if c.column==6: c.number_format="0.0%"; c.font=Font(name=F,size=10,bold=True)
for i,w in enumerate([14,30,12,11,12,12,10,70,26]): ws.column_dimensions[chr(65+i)].width=w
ws.row_dimensions[r0].height=42; ws.freeze_panes="C5"
n=r+2
ws.cell(n,1,"Вывод: образец — вулканизующая паста для тиоколовых герметиков (тип пасты №9: диоксид марганца в дибутилфталате), а не тиоколовая основа.").font=Font(name=F,bold=True,size=10)
ws.cell(n+1,1,"Итог = среднее по ИК и ЯМР, приведённое к сумме 100 %. До получения ЯМР итог равен оценке по ИК. Не обнаружено — содержание ниже предела обнаружения ИК (≈2–5 %).").font=Font(name=F,italic=True,size=9)
ns=wb.create_sheet("Примечания")
notes=["Примечания к оценке состава",
"1. База пересчёта — 100 % органической части. Неорганический компонент исключён; его массовую долю ИК-методом определить нельзя (нужна зольность), природу металла — РФА.",
"2. ИК: НПВО, алмазный кристалл, 4000–560 см⁻¹. Обработка та же, что для образцов buryak_1t, 2t, 1p, 2x; для проверки на тиокол спектр сопоставлен с чистым тиоколом buryak_1t.",
"3. Пределы обнаружения ИК: тиокол ≈5 % (по полосе 1024 см⁻¹ на фоне полос ДБФ), прочие органические компоненты ≈2–3 %.",
"4. ЯМР ¹H и ¹³C будет проведён дополнительно: подтверждение ДБФ (4,30 т, 7,5/7,7 м.д.), проверка на диизобутилфталат, хлорпарафин и стеариновую кислоту (2,35 т, 1,25 м.д.)."]
for i,t in enumerate(notes,1):
    c=ns.cell(i,1,t); c.font=Font(name=F,bold=(i==1),size=12 if i==1 else 10); c.alignment=Alignment(wrap_text=True,vertical="top")
ns.column_dimensions["A"].width=120
for sh in (ws,ns):
    sh.page_setup.orientation="landscape"; sh.page_setup.paperSize=9; sh.page_setup.fitToWidth=1; sh.page_setup.fitToHeight=0; sh.sheet_properties.pageSetUpPr.fitToPage=True
wb.save(os.path.join(HERE,"Состав_buryak_shmurdyak_black_ИК.xlsx"))
