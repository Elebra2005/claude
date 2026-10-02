"""Смета обвязки в Excel из JSON: лист «По узлам схемы» + лист «Сводная к заказу» (одинаковые позиции сведены, по поставщикам).

Запуск:  python3 smeta_xlsx.py smeta-data.json out.xlsx

Формат JSON:
{"groups": [{"title": "Узел …", "rows": [
   {"pos": "V-P3", "name": "Кран шаровой кламп …", "size": "К25 (½″)", "purpose": "где стоит",
    "qty_scheme": "1", "qty_buy": "2", "status": "КУПИТЬ|УТОЧНИТЬ|ЕСТЬ",
    "shop": "Гросснер: кран кламп", "url": "https://…?attribute_pa_…", "extra_url": null}]}]}
Количества — строками («1», «≈3 м»): числа складываются в сводной, остальное склеивается через « + ».
Поставщик в сводной — текст `shop` до двоеточия.
"""
import json, math, re, sys
from urllib.parse import urlparse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L

NAVY = '1F3864'
HF = PatternFill('solid', fgColor=NAVY); SF = PatternFill('solid', fgColor='D9E2F3'); ZF = PatternFill('solid', fgColor='F7F9FC')
th = Side(style='thin', color='BFBFBF'); B = Border(left=th, right=th, top=th, bottom=th)
SB = Border(left=th, right=th, top=Side(style='medium', color=NAVY), bottom=th)
W = Alignment(wrap_text=True, vertical='top'); C = Alignment(wrap_text=True, vertical='top', horizontal='center')
STAT = {'КУПИТЬ': ('E2EFDA', '375623'), 'УТОЧНИТЬ': ('FFF2CC', '7F6000'), 'ЕСТЬ': ('EDEDED', '404040')}


def f(sz=10, b=False, c='000000', u=None):
    return Font(name='Arial', size=sz, bold=b, color=c, underline=u)


def lines(v, w):
    if v is None:
        return 1
    return sum(max(1, math.ceil(len(p) * 1.1 / max(w - 1, 1))) for p in str(v).split('\n'))


def write(ws, hdr, wid, cen, sections, link_col, extra_col=None, stat_col=None, bold_col=None):
    ws.sheet_view.showGridLines = False
    for i, h in enumerate(hdr, 1):
        c = ws.cell(1, i, h); c.font = f(10, True, 'FFFFFF'); c.fill = HF; c.border = B
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
    ws.row_dimensions[1].height = 30
    r = 2; k = 0
    for title, rows in sections:
        c = ws.cell(r, 1, title); c.font = f(10.5, True, NAVY); c.alignment = Alignment(vertical='center')
        for i in range(1, len(hdr) + 1):
            ws.cell(r, i).fill = SF; ws.cell(r, i).border = SB
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(hdr)); ws.row_dimensions[r].height = 20; r += 1
        zebra = False
        for vals, link, extra in rows:
            k += 1; vals = [k] + list(vals); hmax = 1
            for i, v in enumerate(vals, 1):
                v = None if v == '' else v
                c = ws.cell(r, i, v); c.border = B; c.alignment = C if i in cen else W; c.font = f(10)
                if zebra:
                    c.fill = ZF
                hmax = max(hmax, lines(v, wid[i - 1]))
            ws.cell(r, 1).font = f(9, c='7F7F7F')
            if bold_col:
                ws.cell(r, bold_col).font = f(10, True)
            if link:
                ws.cell(r, link_col).hyperlink = link; ws.cell(r, link_col).font = f(10, c='0563C1', u='single')
            if extra_col and extra:
                e = ws.cell(r, extra_col, urlparse(extra).netloc.replace('www.', '')); e.hyperlink = extra; e.font = f(10, c='0563C1', u='single')
            if stat_col and vals[stat_col - 1] in STAT:
                bg, fg = STAT[vals[stat_col - 1]]
                ws.cell(r, stat_col).fill = PatternFill('solid', fgColor=bg); ws.cell(r, stat_col).font = f(9, True, fg)
            ws.row_dimensions[r].height = 13.5 * hmax + 3; zebra = not zebra; r += 1
    for i, w in enumerate(wid, 1):
        ws.column_dimensions[L(i)].width = w
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = f'A1:{L(len(hdr))}{r - 1}'
    ws.page_setup.orientation = 'landscape'; ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0; ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = '1:1'; ws.page_margins.left = ws.page_margins.right = 0.4
    ws.oddFooter.center.text = '&8&A — стр. &P из &N'
    return k


def num(x):
    try:
        return float(str(x).replace(',', '.'))
    except ValueError:
        return None


def total(vs):
    if len(vs) == 1:
        return vs[0]
    ns = [num(v) for v in vs]
    if all(x is not None for x in ns):
        s = sum(ns); return str(int(s)) if s == int(s) else str(s).replace('.', ',')
    return ' + '.join(vs)


def norm(name):
    return re.sub(r'\s*\(≈[^)]*₽\)', '', name).strip().rstrip(',')


def build(data, out):
    groups = data['groups']
    wb = Workbook()
    s1 = [(g['title'], [((x['pos'], x['name'], x['size'], x['purpose'], x['qty_scheme'], x['qty_buy'], x['status'], x['shop']),
                         x.get('url'), x.get('extra_url')) for x in g['rows']]) for g in groups]
    n1 = write(wb.active, ['№', 'Поз. на схеме', 'Наименование', 'Размер', 'Назначение / где стоит', 'По схеме', 'К закупке', 'Статус', 'Где купить', 'Доп. ссылка'],
               [5, 11, 44, 15, 50, 10, 10, 11, 30, 16], {1, 2, 4, 6, 7, 8}, s1, 9, 10, 8, 7)
    wb.active.title = 'По узлам схемы'
    # сводная: одинаковые (наименование, размер) — одной строкой, по поставщикам
    agg, order = {}, []
    for g in groups:
        node = g['title'].split(':')[0]
        for x in g['rows']:
            if x['status'] == 'ЕСТЬ':
                continue
            key = (norm(x['name']), x['size'])
            if key not in agg:
                agg[key] = dict(name=key[0], size=x['size'], n=[], q=[], where=[], st=set(), shop=x['shop'], url=x.get('url'),
                                price=re.search(r'≈[^)]*₽', x['name']))
                order.append(key)
            a = agg[key]; a['n'].append(x['qty_scheme']); a['q'].append(x['qty_buy']); a['st'].add(x['status'])
            a['where'].append((x['pos'] + ': ' if x['pos'] else '') + node)
    sup = {}
    for key in order:
        a = agg[key]; sup.setdefault(a['shop'].split(':')[0].strip(), []).append(a)
    s2 = []
    for sp in sorted(sup):
        rows = []
        for a in sorted(sup[sp], key=lambda a: a['name']):
            nm = a['name'] + (f" ({a['price'].group(0)})" if a['price'] else '')
            st = 'УТОЧНИТЬ' if 'УТОЧНИТЬ' in a['st'] else ('КУПИТЬ' if 'КУПИТЬ' in a['st'] else '/'.join(sorted(a['st'])))
            rows.append(((nm, a['size'], '; '.join(dict.fromkeys(a['where'])), total(a['n']), total(a['q']), st, a['shop']), a['url'], None))
        s2.append((sp, rows))
    n2 = write(wb.create_sheet('Сводная к заказу'), ['№', 'Наименование', 'Размер', 'Где стоит (узлы)', 'По схеме', 'К закупке', 'Статус', 'Где купить'],
               [5, 52, 16, 48, 11, 11, 11, 32], {1, 3, 5, 6, 7}, s2, 8, None, 7, 6)
    wb.save(out)
    print(f'{out}: по узлам {n1} строк, сводная {n2} строк')


if __name__ == '__main__':
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    build(json.load(open(sys.argv[1], encoding='utf-8')), sys.argv[2])
