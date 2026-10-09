"""Собирает Word-файл курса из content_*.py и схем в diagrams/."""
import os, math, datetime
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

import content_m1_m2, content_m3_m4, content_m5_m6

HERE = os.path.dirname(os.path.abspath(__file__))
POSTS = content_m1_m2.POSTS + content_m3_m4.POSTS + content_m5_m6.POSTS
assert [p["n"] for p in POSTS] == list(range(1, 33))

MODULES = {
    1: "Основы: что такое ИИ и как он работает",
    2: "Промтинг: как ставить задачи",
    3: "Возможности Claude",
    4: "Code в деле",
    5: "Автоматизация",
    6: "«Claude-завод»: команда агентов",
}
INK = RGBColor(0x1F, 0x1D, 0x1A)
MUTED = RGBColor(0x6B, 0x66, 0x5E)
ACCENT = RGBColor(0xC2, 0x60, 0x3E)
TEAL = RGBColor(0x2F, 0x6F, 0x6A)
FONT = "Calibri"


def week(n):
    return math.ceil(n / 2)


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def borders(cell, left=None, others="E2DBCF", left_sz=24):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{side}")
        if side == "left" and left:
            e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(left_sz)); e.set(qn("w:color"), left)
        else:
            e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), others)
        b.append(e)
    tcPr.append(b)


def cell_margins(cell, top=120, bottom=120, left=200, right=200):
    tcPr = cell._tc.get_or_add_tcPr()
    m = OxmlElement("w:tcMar")
    for k, v in (("top", top), ("bottom", bottom), ("start", left), ("end", right)):
        e = OxmlElement(f"w:{k}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); m.append(e)
    tcPr.append(m)


def para(container, text="", size=11, bold=False, italic=False, color=INK, after=4, before=0, align=None, keep=False):
    p = container.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(after); pf.space_before = Pt(before); pf.line_spacing = 1.15
    if keep: pf.keep_with_next = True
    if align: p.alignment = align
    if text:
        r = p.add_run(text); r.font.size = Pt(size); r.bold = bold; r.italic = italic; r.font.color.rgb = color; r.font.name = FONT
    return p


def run(p, text, size=11, bold=False, italic=False, color=INK):
    r = p.add_run(text); r.font.size = Pt(size); r.bold = bold; r.italic = italic; r.font.color.rgb = color; r.font.name = FONT
    return r


def label(doc, text, color=ACCENT):
    return para(doc, text.upper(), size=9, bold=True, color=color, after=3, before=10, keep=True)


def fix_widths(t, widths):
    t.autofit = False
    tblPr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    grid = t._tbl.tblGrid
    for j, gc in enumerate(grid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int(widths[j].twips)))
    for row in t.rows:
        for j, c in enumerate(row.cells):
            c.width = widths[j]


def text_box(doc, text, fill="F7F4EF", left="C2603E"):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.cell(0, 0)
    c.width = Cm(16.5)
    shade(c, fill); borders(c, left=left); cell_margins(c, 180, 180, 260, 220)
    fix_widths(t, [Cm(16.5)])
    c.paragraphs[0].text = ""
    first = True
    for line in text.split("\n"):
        p = c.paragraphs[0] if first else c.add_paragraph()
        first = False
        pf = p.paragraph_format; pf.space_after = Pt(0); pf.line_spacing = 1.15
        if line.strip() == "":
            pf.space_after = Pt(2)
            continue
        is_title = line.startswith("Урок ") and p is c.paragraphs[0]
        run(p, line, size=11, bold=is_title)
    return t


def table(doc, header, rows, widths, header_fill="1F1D1A", zebra="F7F4EF", size=9.5):
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(header):
        c = t.cell(0, j); c.width = widths[j]
        shade(c, header_fill); borders(c); cell_margins(c, 80, 80, 120, 120)
        p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
        run(p, h, size=size, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    for i, row in enumerate(rows, 1):
        for j, val in enumerate(row):
            c = t.cell(i, j); c.width = widths[j]
            if i % 2 == 0: shade(c, zebra)
            borders(c); cell_margins(c, 70, 70, 120, 120)
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.1
            run(p, str(val), size=size)
    fix_widths(t, widths)
    # повтор шапки на новой странице
    trPr = t.rows[0]._tr.get_or_add_trPr(); h = OxmlElement("w:tblHeader"); h.set(qn("w:val"), "true"); trPr.append(h)
    return t


def bullets(doc, items, size=10.5):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        if isinstance(it, tuple):
            run(p, it[0] + ": ", size=size, bold=True); run(p, it[1], size=size)
        else:
            run(p, it, size=size)


def media_word(kind):
    return {"Схема": "схема", "Видео": "видео", "Скриншот": "скриншот", "Опрос": "опрос",
            "Документ": "документ", "Карточка": "карточка"}.get(kind, kind.lower())


def build(out):
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(2.0); sec.top_margin = Cm(1.8); sec.bottom_margin = Cm(1.8)
    st = doc.styles["Normal"]; st.font.name = FONT; st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    for name, size, color in (("Heading 1", 20, INK), ("Heading 2", 15, INK), ("Heading 3", 12, ACCENT)):
        s = doc.styles[name]; s.font.name = FONT; s.font.size = Pt(size); s.font.bold = True; s.font.color.rgb = color
        s.element.rPr.rFonts.set(qn("w:asciiTheme"), "") if False else None
        rf = s.element.rPr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts"); s.element.rPr.append(rf)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"): rf.set(qn(a), FONT)
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]

    # ---- Обложка
    para(doc, "", after=110)
    para(doc, "КУРС МИНИ-УРОКОВ ДЛЯ РАБОЧЕГО КАНАЛА", size=11, bold=True, color=ACCENT, after=8)
    para(doc, "ИИ в работе: от основ до команды агентов", size=28, bold=True, after=10)
    para(doc, "32 обучающих поста о работе с Claude: готовые тексты, схемы, сценарии видео и пометки, какие скриншоты и медиа приложить.", size=13, color=MUTED, after=30)
    for m, name in MODULES.items():
        ns = [p["n"] for p in POSTS if p["module"] == m]
        para(doc, f"Модуль {m}. {name} — посты {ns[0]}–{ns[-1]}", size=11.5, after=4)
    para(doc, "", after=40)
    para(doc, f"Версия от {datetime.date.today().strftime('%d.%m.%Y')}", size=10, color=MUTED)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---- Как пользоваться
    doc.add_heading("Как пользоваться документом", level=1)
    para(doc, "Каждый пост ниже готов к публикации: текст можно копировать в канал целиком. Под текстом указано, какие медиа приложить, а для постов с видео — пошаговый сценарий записи экрана.", after=8)
    doc.add_heading("Ритм публикации", level=3)
    para(doc, "Два поста в неделю, например во вторник и четверг. Весь курс занимает 16 недель. Если коллеги не успевают, лучше растянуть курс, чем сократить посты.", after=6)
    doc.add_heading("Структура каждого поста", level=3)
    bullets(doc, [("Заголовок", "«Урок N. Тема» — чтобы посты было легко искать в канале"),
                  ("Суть", "3–5 коротких абзацев или списков без лишней теории"),
                  ("Пример", "промт или ситуация из нашей работы"),
                  ("Задание «Попробуйте»", "одно действие на 5 минут — закрепляет навык"),
                  ("Медиа", "схема, скриншот или видео")])
    doc.add_heading("Обозначения медиа", level=3)
    table(doc, ["Тип", "Что это и кто готовит"], [
        ["Схема", "Готовая картинка, вставлена в документ под постом. Её можно сохранить из Word (правой кнопкой мыши → «Сохранить как рисунок») или взять из папки со схемами."],
        ["Скриншот", "Снимок экрана из вашего интерфейса Claude. Делается перед публикацией, личные и рабочие данные замазываются."],
        ["Видео", "Запись экрана 1–3 минуты (в финальном уроке — 10–15 минут). Сценарий — в таблице под постом."],
        ["Опрос", "Встроенный опрос мессенджера. Варианты ответов указаны под постом."],
        ["Карточка / документ", "Готовая картинка-памятка или повторная отправка мануала."],
    ], [Cm(3.2), Cm(13.6)])
    doc.add_heading("Советы по записи видео", level=3)
    bullets(doc, [
        "Одно видео — одна мысль. Длина 1–3 минуты, кроме финального кейса.",
        "Записывайте с голосом, но в тексте поста коротко пишите, что показано: многие смотрят без звука.",
        "Перед записью закройте лишние вкладки и уведомления, используйте учебные или обезличенные данные.",
        "Сделайте пробный прогон: ответы Claude каждый раз немного отличаются, и лучше заранее знать, что он покажет.",
        "Долгое ожидание ответа ускоряйте при монтаже или вырезайте.",
        "Время в сценариях ориентировочное: это длительность фрагмента после монтажа.",
        "Названия кнопок и пунктов меню в интерфейсе могут меняться с обновлениями. Перед записью сверьте их со своей версией приложения.",
    ])
    doc.add_heading("Постоянные рубрики между уроками", level=3)
    bullets(doc, [("Промт недели", "готовый промт под типовую задачу компании"),
                  ("Вопрос из канала", "разбор вопроса коллеги (с его согласия)"),
                  ("Ошибка недели", "где ИИ ошибся и как это поймали — учит привычке проверять"),
                  ("Опрос раз в месяц", "что непонятно и какие темы нужны — помогает корректировать план")])
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---- Контент-план
    doc.add_heading("Контент-план", level=1)
    rows = []
    for p in POSTS:
        kinds = []
        for k, _ in p["media"]:
            w = media_word(k)
            if w not in kinds: kinds.append(w)
        rows.append([p["n"], week(p["n"]), p["title"], ", ".join(kinds)])
    table(doc, ["№", "Неделя", "Тема", "Медиа"], rows, [Cm(1.0), Cm(1.6), Cm(10.2), Cm(4.0)])

    # ---- Посты
    current = None
    for p in POSTS:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        if p["module"] != current:
            current = p["module"]
            para(doc, f"МОДУЛЬ {current}", size=10, bold=True, color=ACCENT, after=0)
            doc.add_heading(MODULES[current], level=1)
        h = doc.add_heading(f"Пост {p['n']}. {p['title']}", level=2)
        kinds = []
        for k, _ in p["media"]:
            w = media_word(k)
            if w not in kinds: kinds.append(w)
        meta = para(doc, "", after=2)
        run(meta, f"Модуль {p['module']} · неделя {week(p['n'])} · медиа: {', '.join(kinds)}", size=9.5, color=MUTED)
        g = para(doc, "", after=6)
        run(g, "Цель урока: ", size=10.5, bold=True); run(g, p["goal"], size=10.5)

        label(doc, "Текст поста — готов к публикации")
        text_box(doc, p["text"])

        ml = label(doc, "Медиа к посту", color=TEAL)
        bullets(doc, [(k, d) for k, d in p["media"]])
        if p.get("diagram"):
            # держим подпись «Медиа» и пункты вместе со схемой
            for q in doc.paragraphs[-(len(p["media"]) + 1):]:
                q.paragraph_format.keep_with_next = True
            img = os.path.join(HERE, "diagrams", p["diagram"])
            pi = doc.add_paragraph(); pi.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pi.paragraph_format.space_before = Pt(6); pi.paragraph_format.keep_with_next = True
            pi.add_run().add_picture(img, width=Cm(15.5))
            para(doc, f"Схема к посту {p['n']} (файл {p['diagram']})", size=9, italic=True, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)

        if p.get("video"):
            total = sum(s for _, _, s in p["video"])
            mins, secs = divmod(total, 60)
            dur = f"{mins} мин {secs:02d} с" if mins else f"{secs} с"
            label(doc, f"Сценарий видео — около {dur}", color=TEAL)
            rows = [[i + 1, scr, say, f"{sec} с"] for i, (scr, say, sec) in enumerate(p["video"])]
            table(doc, ["№", "Что на экране", "Что говорить", "Время"], rows,
                  [Cm(0.9), Cm(7.6), Cm(6.6), Cm(1.7)], header_fill="2F6F6A", zebra="EEF5F3")

        if p.get("note"):
            label(doc, "Заметка для автора (не публикуется)", color=MUTED)
            para(doc, p["note"], size=10, italic=True, color=MUTED)

    doc.save(out)
    print("saved", out)


if __name__ == "__main__":
    build(os.path.join(os.path.dirname(HERE), "Курс_ИИ_в_работе_32_поста.docx"))
